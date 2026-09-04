"""
Load the Completely_Sold symbol list from the canonical BuySell Excel workbook.

Returns {SYMBOL: {"avg_buy": float, "avg_sell": float}} for each completely-sold
position so AlertApp can show Current / Avg Buy / Avg Sold and alert when price
drops ≥ drop_pct below avg_sell (rebuy candidate).

Default workbook: C:\\Investment\\reports\\IBKR_BuySell_Since_2020.xlsx
Override via workbook_path arg or AlertApp config.ini [rebuy] workbook_path.

Completely_Sold sheet columns (0-based):
  0 Symbol, 1 Buy_Qty_Total, 2 Sell_Qty_Total,
  5 Total_Buy_Cost, 6 Total_Sell_Proceeds, 10 Last_Sold_Price
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, TypedDict


class SoldLevels(TypedDict):
    avg_buy: float
    avg_sell: float


_DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[1] / "reports" / "IBKR_BuySell_Since_2020.xlsx"
)
_SHEET_NAME = "Completely_Sold"

_COL_SYMBOL = 0
_COL_BUY_QTY = 1
_COL_SELL_QTY = 2
_COL_BUY_COST = 5
_COL_SELL_PROCEEDS = 6
_COL_LAST_SOLD_PRICE = 10


def _safe_avg(total: object, qty: object) -> float | None:
    try:
        q = float(qty)  # type: ignore[arg-type]
        t = float(total)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    if q <= 0:
        return None
    return t / q


def load_sold_watchlist(
    workbook_path: Optional[Path] = None,
    *,
    exclude_symbols: Optional[list[str]] = None,
) -> dict[str, SoldLevels]:
    """
    Return {SYMBOL: {"avg_buy": ..., "avg_sell": ...}} from Completely_Sold.

    Empty dict if the workbook cannot be read (main watchlist still works).
    """
    path = workbook_path or _DEFAULT_WORKBOOK
    skip = {s.upper() for s in (exclude_symbols or [])}

    if not Path(path).is_file():
        print(
            f"[sold_watchlist] Workbook not found: {path}. "
            "Rebuy watchlist will be empty.",
            flush=True,
        )
        return {}

    try:
        import openpyxl
    except ImportError:
        print(
            "[sold_watchlist] openpyxl not installed; rebuy watchlist disabled. "
            "Run: pip install openpyxl",
            flush=True,
        )
        return {}

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        print(f"[sold_watchlist] Cannot open workbook {path}: {exc}", flush=True)
        return {}

    if _SHEET_NAME not in wb.sheetnames:
        print(
            f"[sold_watchlist] Sheet '{_SHEET_NAME}' not found in {path}. "
            f"Available: {wb.sheetnames}",
            flush=True,
        )
        wb.close()
        return {}

    ws = wb[_SHEET_NAME]
    result: dict[str, SoldLevels] = {}
    first_row = True

    for row in ws.iter_rows(values_only=True):
        if first_row:
            first_row = False
            continue

        symbol = row[_COL_SYMBOL]
        if not symbol:
            continue
        symbol = str(symbol).strip().upper()
        if symbol in skip:
            continue

        avg_sell = _safe_avg(row[_COL_SELL_PROCEEDS], row[_COL_SELL_QTY])
        if avg_sell is None and row[_COL_LAST_SOLD_PRICE] is not None:
            try:
                avg_sell = float(row[_COL_LAST_SOLD_PRICE])
            except (TypeError, ValueError):
                avg_sell = None

        avg_buy = _safe_avg(row[_COL_BUY_COST], row[_COL_BUY_QTY])

        if avg_sell is None or avg_sell <= 0:
            continue
        if avg_buy is None or avg_buy <= 0:
            # Still track sell-only if buy cost missing (rare)
            avg_buy = avg_sell

        result[symbol] = {
            "avg_buy": round(avg_buy, 4),
            "avg_sell": round(avg_sell, 4),
        }

    wb.close()
    print(
        f"[sold_watchlist] Loaded {len(result)} completely-sold symbols "
        f"from {path.name}.",
        flush=True,
    )
    return result
