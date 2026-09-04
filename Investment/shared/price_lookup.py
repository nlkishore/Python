"""
On-demand symbol lookup: weighted-average sold price from the BuySell Excel
Sells sheet + last Yahoo market price. Used by AlertApp PRICE / Q commands.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Mapping, Optional

_DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[1] / "reports" / "IBKR_BuySell_Since_2020.xlsx"
)
_SELLS_SHEET = "Sells"

_HEADER_ALIASES = {
    "symbol": "symbol",
    "date": "date",
    "quantity": "quantity",
    "qty": "quantity",
    "price": "price",
    "net amount": "net",
    "net_amount": "net",
    "gross amount": "gross",
    "gross_amount": "gross",
    "transaction type": "side",
    "buy/sell": "side",
}


@dataclass(frozen=True)
class SellStats:
    avg_sell: float
    sell_qty: float
    proceeds: float
    last_sold: Optional[date]


_cache_key: tuple[str, float] | None = None
_cache_data: dict[str, SellStats] = {}


def clear_sell_averages_cache() -> None:
    global _cache_key, _cache_data
    _cache_key = None
    _cache_data = {}


def _as_float(value: object) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%b-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10] if fmt.startswith("%Y") else text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "")[:19]).date()
    except ValueError:
        return None


def _row_proceeds(qty: float, price: float | None, net: float | None, gross: float | None) -> float:
    if net is not None and abs(net) > 1e-12:
        return abs(net)
    if gross is not None and abs(gross) > 1e-12:
        return abs(gross)
    if price is not None and qty > 0:
        return abs(price) * qty
    return 0.0


def aggregate_sell_rows(rows: Iterable[Mapping[str, object]]) -> dict[str, SellStats]:
    """
    Build per-symbol weighted average sell price from already-normalized rows.

    Each row may include: symbol, quantity, price, net, gross, date, side.
    Rows with side set to something other than sell are skipped.
    Avg sold = total proceeds / total qty (proceeds prefer Net Amount).
    """
    buckets: dict[str, list[tuple[float, float, date | None]]] = {}
    for raw in rows:
        side = str(raw.get("side") or "SELL").strip().upper()
        if side and side not in ("SELL", "S", "SLD"):
            continue
        symbol = str(raw.get("symbol") or "").strip().upper()
        if not symbol or symbol in ("-", "NAN", "NONE"):
            continue
        qty = _as_float(raw.get("quantity"))
        if qty is None or qty <= 1e-12:
            continue
        qty = abs(qty)
        proceeds = _row_proceeds(
            qty,
            _as_float(raw.get("price")),
            _as_float(raw.get("net")),
            _as_float(raw.get("gross")),
        )
        if proceeds <= 1e-12:
            continue
        buckets.setdefault(symbol, []).append((qty, proceeds, _as_date(raw.get("date"))))

    out: dict[str, SellStats] = {}
    for symbol, parts in buckets.items():
        qty_sum = sum(q for q, _, _ in parts)
        proc_sum = sum(p for _, p, _ in parts)
        if qty_sum <= 1e-12:
            continue
        dates = [d for _, _, d in parts if d is not None]
        out[symbol] = SellStats(
            avg_sell=round(proc_sum / qty_sum, 4),
            sell_qty=round(qty_sum, 4),
            proceeds=round(proc_sum, 2),
            last_sold=max(dates) if dates else None,
        )
    return out


def _header_map(headers: tuple[object, ...]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for idx, raw in enumerate(headers):
        key = str(raw or "").strip().lower()
        alias = _HEADER_ALIASES.get(key)
        if alias:
            mapped[alias] = idx
    return mapped


def _rows_from_worksheet(ws) -> list[dict[str, object]]:
    iterator = ws.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        return []
    cols = _header_map(tuple(header))
    if "symbol" not in cols or "quantity" not in cols:
        return []
    rows: list[dict[str, object]] = []
    for values in iterator:
        if not values:
            continue
        row: dict[str, object] = {
            "symbol": values[cols["symbol"]] if cols["symbol"] < len(values) else None,
            "quantity": values[cols["quantity"]] if cols["quantity"] < len(values) else None,
        }
        if "price" in cols and cols["price"] < len(values):
            row["price"] = values[cols["price"]]
        if "net" in cols and cols["net"] < len(values):
            row["net"] = values[cols["net"]]
        if "gross" in cols and cols["gross"] < len(values):
            row["gross"] = values[cols["gross"]]
        if "date" in cols and cols["date"] < len(values):
            row["date"] = values[cols["date"]]
        if "side" in cols and cols["side"] < len(values):
            row["side"] = values[cols["side"]]
        rows.append(row)
    return rows


def load_sell_averages(
    workbook_path: Optional[Path] = None,
    *,
    force: bool = False,
) -> dict[str, SellStats]:
    """Read the Sells sheet; cache until the file mtime changes or force=True."""
    global _cache_key, _cache_data
    path = Path(workbook_path) if workbook_path else _DEFAULT_WORKBOOK
    if not path.is_file():
        print(
            f"[price_lookup] Workbook not found: {path}. Avg-sold lookup will be empty.",
            flush=True,
        )
        return {}

    mtime = path.stat().st_mtime
    key = (str(path.resolve()), mtime)
    if not force and _cache_key == key:
        return _cache_data

    try:
        import openpyxl
    except ImportError:
        print(
            "[price_lookup] openpyxl not installed; avg-sold lookup disabled. "
            "Run: pip install openpyxl",
            flush=True,
        )
        return {}

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        print(f"[price_lookup] Cannot open workbook {path}: {exc}", flush=True)
        return {}

    if _SELLS_SHEET not in wb.sheetnames:
        print(
            f"[price_lookup] Sheet '{_SELLS_SHEET}' not found in {path.name}. "
            f"Available: {wb.sheetnames}",
            flush=True,
        )
        wb.close()
        return {}

    try:
        data = aggregate_sell_rows(_rows_from_worksheet(wb[_SELLS_SHEET]))
    finally:
        wb.close()

    _cache_key = key
    _cache_data = data
    print(
        f"[price_lookup] Loaded avg-sold for {len(data)} symbols from {path.name}.",
        flush=True,
    )
    return data


def fetch_yahoo_last(symbol: str) -> tuple[Optional[float], Optional[str]]:
    """
    Return (last_price, as_of_label). Uses 1-day Yahoo history Close.
    After hours this is the last regular-session print, not a live bid/ask.
    """
    ticker = str(symbol or "").strip().upper()
    if not ticker:
        return None, None
    try:
        import yfinance as yf
    except ImportError:
        return None, "yfinance not installed"

    try:
        hist = yf.Ticker(ticker).history(period="1d")
    except Exception as exc:  # noqa: BLE001
        print(f"[price_lookup] Yahoo fetch error for {ticker}: {exc}", flush=True)
        return None, None
    if hist is None or hist.empty or "Close" not in hist.columns:
        return None, None
    price = float(hist["Close"].iloc[-1])
    as_of = None
    try:
        idx = hist.index[-1]
        as_of = pd_timestamp_label(idx)
    except Exception:  # noqa: BLE001
        as_of = "Yahoo last"
    return price, as_of or "Yahoo last"


def pd_timestamp_label(idx: object) -> str:
    if hasattr(idx, "strftime"):
        try:
            hour = getattr(idx, "hour", None)
            minute = getattr(idx, "minute", None)
            if hour == 0 and minute == 0:
                return f"Yahoo last {idx.strftime('%Y-%m-%d')}"
            return f"Yahoo last {idx.strftime('%Y-%m-%d %H:%M')}"
        except (TypeError, ValueError):
            pass
    return "Yahoo last"


def parse_price_command(msg_text: str) -> str | None:
    """
    Return the ticker for PRICE / Q commands, or None if the message is not one.

    Empty string means the verb was present but no symbol was given.
    """
    text = (msg_text or "").strip().upper()
    if not text:
        return None
    for verb in ("PRICE", "Q"):
        if text == verb:
            return ""
        prefix = verb + " "
        if text.startswith(prefix):
            rest = text[len(prefix) :].strip()
            if not rest:
                return ""
            return rest.split()[0]
    return None


def format_price_reply(
    symbol: str,
    stats: Optional[SellStats],
    market: Optional[float],
    as_of: Optional[str] = None,
) -> str:
    """WhatsApp body for one symbol (avg sold + market)."""
    sym = str(symbol or "").strip().upper()
    lines = [f"*{sym}*"]

    if stats is None:
        lines.append("Avg Sold:    (no sells in BuySell workbook)")
    else:
        qty = stats.sell_qty
        qty_txt = f"{qty:.4f}".rstrip("0").rstrip(".")
        lines.append(f"Avg Sold:    ${stats.avg_sell:,.2f}   ({qty_txt} sh)")

    if market is None:
        note = as_of or "no Yahoo quote — delisted or unknown ticker"
        lines.append(f"Market:      ({note})")
    else:
        label = as_of or "Yahoo last"
        lines.append(f"Market:      ${market:,.2f}   ({label})")

    if stats is not None and market is not None and stats.avg_sell > 1e-12:
        vs_pct = (market - stats.avg_sell) / stats.avg_sell * 100.0
        lines.append(f"vs Avg Sold: {vs_pct:+.1f}%")

    if stats is not None and stats.last_sold is not None:
        lines.append(f"Last sold:   {stats.last_sold.isoformat()}")

    return "\n".join(lines)


def lookup_price_reply(symbol: str, workbook_path: Optional[Path] = None) -> str:
    """Load avg sold + Yahoo last and format the WhatsApp reply."""
    ticker = str(symbol or "").strip().upper()
    stats_map = load_sell_averages(workbook_path)
    stats = stats_map.get(ticker)
    market, as_of = fetch_yahoo_last(ticker)
    return format_price_reply(ticker, stats, market, as_of)
